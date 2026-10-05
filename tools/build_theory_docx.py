#!/usr/bin/env python3
"""Build source-derived Vietnamese theory reference; do not modify the template."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re
import subprocess
import sys
from lxml import etree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = Path('/tmp/agv_theory_20261005')
ASSET = ROOT / 'docs/theory_report'
REF = ROOT / 'docs/PSTMO.docx'
OUT = ROOT / 'docs/CO_SO_LY_THUYET_TOAN_DIEN_DU_AN_AGV_PSTMO.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
NS = {'w': W, 'm': M}

def sha(data):
    return hashlib.sha256(data).hexdigest()

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e

def audit_template():
    with ZipFile(REF) as z:
        inventory = {n: {'bytes': len(z.read(n)), 'sha256': sha(z.read(n))} for n in z.namelist()}
        root = ET.fromstring(z.read('word/document.xml'))
        baseline = {'reference': str(REF), 'sha256': sha(REF.read_bytes()),
                    'rendered_pages': len(list((WORK/'reference').glob('page-*.png'))),
                    'sections': [ET.tostring(x).decode() for x in root.findall('.//w:sectPr', NS)],
                    'parts': inventory}
    (WORK/'baseline.json').write_text(json.dumps(baseline, indent=2), encoding='utf-8')
    return baseline

def make_review_sheets(folder, prefix, cols=5, rows=4):
    files = sorted(folder.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1]))
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 17)
    for start in range(0, len(files), cols*rows):
        canvas = Image.new('RGB',(cols*300,rows*443),'#d2d2d2')
        draw = ImageDraw.Draw(canvas)
        for idx, f in enumerate(files[start:start+cols*rows]):
            pic = Image.open(f).convert('RGB'); pic.thumbnail((294,416))
            x=(idx%cols)*300; y=(idx//cols)*443
            canvas.paste(pic,(x,y+24)); draw.text((x+7,y+2),f.stem,font=font,fill='black')
        canvas.save(WORK/f'{prefix}_{start//(cols*rows)+1:02d}.jpg',quality=92)

class Builder:
    def __init__(self):
        self.baseline=audit_template()
        # Working package derived from retained reference, never blank Document().
        self.d=Document(REF)
        body=self.d._element.body
        for child in list(body):
            if child.tag != qn('w:sectPr'): body.remove(child)
        self.figs=[]; self.equations=0; self.tables=0; self.headings=[]; self.bm=1
        self.width=(self.d.sections[0].page_width-self.d.sections[0].left_margin-self.d.sections[0].right_margin)/914400
        title=self.d.styles['Title']; src=self.d.styles['Report Title']
        for child in list(title.element):
            if child.tag in [qn('w:rPr'),qn('w:pPr')]: title.element.remove(child)
        for tag in ['w:pPr','w:rPr']:
            orig=src.element.find(qn(tag))
            if orig is not None: title.element.append(deepcopy(orig))
        for border in title.element.findall('.//'+qn('w:pBdr')): border.getparent().remove(border)
        title.font.color.rgb=RGBColor(0,0,0)
        title.font.size=Pt(16)
        props=self.d.core_properties
        props.title='Cơ sở lý thuyết và đối chiếu triển khai dự án AGV PSTMO'
        props.subject='Tài liệu lý thuyết và truy vết mã nguồn ngày 06/10/2026'
        props.author=''; props.last_modified_by=''; props.comments=''
        self.sources=[]

    def p(self, text='', style='Normal'):
        p=self.d.add_paragraph(text,style)
        p.paragraph_format.widow_control=True
        if style=='Normal':
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after=Pt(4)
        return p

    def heading(self,text,level=1):
        p=self.p(text,'Report Heading '+str(level))
        p.paragraph_format.keep_with_next=True
        if level==1:
            p.paragraph_format.page_break_before=text.startswith(('Chương 1 ', 'Phụ lục '))
            p.paragraph_format.space_before=Pt(14)
        pp=p._p.get_or_add_pPr(); outline=el('w:outlineLvl',val=level-1); pp.append(outline)
        anchor='theory_'+str(self.bm)
        p._p.insert(1,el('w:bookmarkStart',id=self.bm,name=anchor))
        p._p.append(el('w:bookmarkEnd',id=self.bm)); self.bm+=1
        if level==1: self.headings.append((text,anchor))
        return p

    def equation(self,text):
        self.equations+=1
        p=self.p('', 'Report Equation'); p.paragraph_format.keep_together=True
        p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7)
        math=OxmlElement('m:oMath')
        def literal(s,parent):
            if not s: return
            r=OxmlElement('m:r'); rp=OxmlElement('m:rPr'); sty=OxmlElement('m:sty'); sty.set(qn('m:val'),'p');rp.append(sty);r.append(rp)
            txt=OxmlElement('m:t'); txt.text=s; txt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(txt);parent.append(r)
        subs=dict(zip('₀₁₂₃₄₅₆₇₈₉₊₋₌ₓᵧₖᵢⱼₘₙₐ꜀ₚₛₜₒₑᵦ','0123456789+-=xykijmnacpst oeb'.replace(' ','')))
        subs.update({'ₙ':'n','ₗ':'l','ᵤ':'u','ᵣ':'r'})
        sups=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ⁱᴸᴿᵀ','0123456789+-iLRT'))
        def run(s,parent):
            s=s.replace('𝓌','w')
            absolute=re.search(r'abs\(([^()]*)\)',s)
            if absolute:
                run(s[:absolute.start()],parent)
                delimiter=OxmlElement('m:d');props=OxmlElement('m:dPr')
                for tag in ['m:begChr','m:endChr']:
                    char=OxmlElement(tag);char.set(qn('m:val'),'|');props.append(char)
                delimiter.append(props);inside=OxmlElement('m:e')
                run(absolute[1],inside);delimiter.append(inside);parent.append(delimiter)
                run(s[absolute.end():],parent)
                return
            idx=0;buffer=''
            while idx<len(s):
                c=s[idx]
                if c in ('\u0302','\u0307') and buffer:
                    base=buffer[-1];literal(buffer[:-1],parent);buffer=''
                    acc=OxmlElement('m:acc');props=OxmlElement('m:accPr')
                    char=OxmlElement('m:chr');char.set(qn('m:val'),c);props.append(char);acc.append(props)
                    e=OxmlElement('m:e');literal(base,e);acc.append(e);parent.append(acc)
                    idx+=1
                    continue
                if c in subs or c in sups or c=='_':
                    if buffer or len(parent):
                        existing=None
                        if buffer:
                            size=2 if buffer[-1] in ('′','″') and len(buffer)>1 else 1
                            base=buffer[-size:];literal(buffer[:-size],parent);buffer=''
                        else:
                            existing=parent[-1];parent.remove(existing)
                        sub='';sup=''
                        while idx<len(s) and (s[idx] in subs or s[idx] in sups or s[idx]=='_'):
                            if s[idx]=='_':
                                idx+=1
                                match=re.match(r'[A-Za-z0-9]+|[Α-Ωα-ω]',s[idx:])
                                if match:sub+=match[0];idx+=len(match[0])
                            elif s[idx] in subs:sub+=subs[s[idx]];idx+=1
                            else:sup+=sups[s[idx]];idx+=1
                        node=OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
                        e=OxmlElement('m:e')
                        if existing is not None:e.append(existing)
                        else:literal(base,e)
                        node.append(e)
                        if sub:
                            part=OxmlElement('m:sub');literal(sub,part);node.append(part)
                        if sup:
                            part=OxmlElement('m:sup');literal(sup,part);node.append(part)
                        parent.append(node)
                        continue
                buffer+=c;idx+=1
            if buffer:literal(buffer,parent)
        last=0
        for match in re.finditer(r'\[frac:(.*?)\|(.*?)\]',text):
            run(text[last:match.start()],math)
            frac=OxmlElement('m:f')
            for tag,s in [('m:num',match[1]),('m:den',match[2])]:
                part=OxmlElement(tag);run(s,part);frac.append(part)
            math.append(frac);last=match.end()
        run(text[last:],math)
        p._p.append(math)
        p.add_run('    ('+str(self.equations)+')').font.size=Pt(9)

    def figure(self,path,caption):
        prefix,rel=path.split('/',1)
        base={'R':ROOT/'docs/robot_3d_report','P':ROOT/'docs/pstmo_bao_cao_toan_dien_assets/figures','W':ROOT/'docs/warehouse_cross_aisles_5_routes/figures'}[prefix]
        full=base/rel
        with Image.open(full) as im:
            width=min(self.width,6.7);height=width*im.height/im.width
            if height>4.2: width*=4.2/height
        p=self.p();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(2)
        pic=p.add_run().add_picture(str(full),width=Inches(width))
        pic._inline.docPr.set('descr',caption)
        n=len(self.figs)+1
        cap=self.p(f'Hình {n}. {caption}', 'Report Caption')
        cap.paragraph_format.keep_with_next=False
        cap.paragraph_format.keep_together=True
        cap.paragraph_format.space_after=Pt(7)
        self.figs.append({'number':n,'caption':caption,'source':str(full.relative_to(ROOT)),'sha256':sha(full.read_bytes())})

    def table(self,rows,widths=None,small=False):
        self.tables+=1
        table=self.d.add_table(rows=0,cols=len(rows[0]))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        if widths is None: widths=[1/len(rows[0])]*len(rows[0])
        for c,width in zip(table.columns,widths): c.width=Inches(self.width*width)
        props=table._tbl.tblPr
        borders=el('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            borders.append(el('w:'+side,val='single',sz=4,color='D9D9D9'))
        props.append(borders)
        margins=el('w:tblCellMar')
        for side in ['top','bottom','left','right']:margins.append(el('w:'+side,w=65,type='dxa'))
        props.append(margins)
        for i,row in enumerate(rows):
            cells=table.add_row().cells
            trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(el('w:cantSplit'))
            if i==0:trpr.append(el('w:tblHeader'))
            for j,(cell,txt) in enumerate(zip(cells,row)):
                cell.width=Inches(self.width*widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1.0
                p._p.get_or_add_pPr().append(el('w:snapToGrid',val='0'))
                p.paragraph_format.keep_with_next=False;p.paragraph_format.widow_control=True
                r=p.add_run(txt);r.font.name='Times New Roman';r.font.size=Pt(8.5 if small else 9)
                if i==0:
                    r.bold=True;cell._tc.get_or_add_tcPr().append(el('w:shd',fill='EEEEEE'))
        self.p().paragraph_format.space_after=Pt(2)

    def trace(self):
        rows=[['Thành phần và cơ sở','Tệp nguồn chính','Chương']]
        entries=[
          ('Hình học, cấu trúc và giới hạn robot','adaptive_pivot_g2/include/adaptive_pivot_g2/types.hpp','2, 4'),
          ('Bézier, đạo hàm, độ cong, cap bánh','adaptive_pivot_g2/src/quintic_transition.cpp','9'),
          ('Hai d, coarse/recovery/refinement alpha','adaptive_pivot_g2/src/hierarchical_shape_search.cpp','9'),
          ('Conditioning và khử dao động','adaptive_pivot_g2/src/path_conditioning.cpp','9'),
          ('LOS có trong thư viện, không bật trong PSTMO mặc định','adaptive_pivot_g2/src/line_of_sight.cpp','9'),
          ('Tìm d thích nghi nhánh legacy','adaptive_pivot_g2/src/adaptive_search.cpp','11'),
          ('Quét vận tốc, time window và pivot time','adaptive_pivot_g2/src/time_parameterization.cpp','10'),
          ('Chi phí ổn định và lựa chọn ứng viên','adaptive_pivot_g2/src/candidate_selection.cpp','10'),
          ('Quy hoạch động trên góc liên tiếp','adaptive_pivot_g2/src/path_optimization.cpp','10'),
          ('Chính sách Hybrid đối xứng','adaptive_pivot_g2/src/hybrid_selection.cpp','11'),
          ('Pipeline thật và bất biến cuối','adaptive_pivot_g2_nav2/src/adaptive_pivot_g2_smoother.cpp','9, 10'),
          ('Quét phần trong và biên footprint','adaptive_pivot_g2_nav2/src/footprint_safety.cpp','6'),
          ('Hai nhánh và fallback Raw','adaptive_pivot_g2_nav2/src/safety_gated_hybrid_smoother.cpp','11'),
          ('Path geometry và tracking metrics','adaptive_pivot_g2_benchmark/adaptive_pivot_g2_benchmark/compare_paths.py','13'),
          ('Metrics góc cua, vận tốc, định vị, clearance','adaptive_pivot_g2_benchmark/adaptive_pivot_g2_benchmark/*metrics.py','13'),
          ('Hợp đồng Path, yaw đầu và runner','adaptive_pivot_g2_benchmark/adaptive_pivot_g2_benchmark/path_contract.py; initial_heading.py; execution_trial.py; execution_matrix.py','12–14'),
          ('Panel chọn planner/smoother/environment','adaptive_pivot_g2_rviz/src/planner_selector_panel.cpp','1, 12'),
          ('Các node, mô hình và thông số','vacuum_robot_gazebo/config; launch; urdf; models; maps; worlds','3–7, 11–12')]
        for row in entries: rows.append(list(row))
        self.table(rows,[.34,.55,.11],True)
        self.p('Đường dẫn trong bảng tính từ src/. Các tên file dài được phép xuống dòng; danh sách hash dưới đây định danh bản nguồn cụ thể.')

    def config(self):
        for name in ['nav2_params.yaml','bridge.yaml','real_robot_profile.yaml']:
            self.heading('Cấu hình '+name.replace('_',' ').replace('.yaml',''),2)
            full=ROOT/'src/vacuum_robot_gazebo/config'/name
            self.p('Nguồn: '+str(full.relative_to(ROOT)), 'Report References')
            rows=[['Dòng','Khóa và giá trị trong nguồn']]
            for i,line in enumerate(full.read_text().splitlines(),1):
                code=line.split('#')[0].rstrip()
                if code.strip(): rows.append([str(i),code.replace('  ','· ')])
            self.table(rows,[.07,.93],True)

    def manifest(self):
        files=sorted(p for p in (ROOT/'src').rglob('*') if p.is_file() and p.suffix in {'.cpp','.hpp','.py','.yaml','.xml','.urdf','.sdf'} and '__pycache__' not in str(p))
        files += [ROOT/'docs/robot_3d_report/README.md',ROOT/'docs/warehouse_cross_aisles_5_routes/README.md',REF]
        files += list((ROOT/'file_3D').glob('*.step'))
        files += list((ROOT/'src/vacuum_robot_gazebo').rglob('*.stl'))
        entries=[{'path':str(p.relative_to(ROOT)),'sha256':sha(p.read_bytes()),'bytes':p.stat().st_size} for p in files]
        report={'date':'2026-10-06','git_head':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
                'files':entries,'scope':'Source snapshot; does not establish binary loaded in an earlier run.'}
        (ASSET/'source_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
        self.p('Git HEAD đối chiếu: '+report['git_head']+'. Hash dưới đây là SHA-256 rút gọn 16 ký tự để tra; bản đầy đủ nằm trong docs/theory_report/source_manifest.json. Snapshot nguồn không thay thế xác nhận binary runtime của những lần chạy cũ.')
        selected=[e for e in entries if ('/src/' in e['path'] and e['path'].endswith('.cpp')) or '/config/' in e['path'] or e['path'].endswith(('.urdf','model.sdf'))]
        self.table([['Tệp nguồn','SHA 256 rút gọn']]+[[e['path'],e['sha256'][:16]] for e in selected],[.76,.24],True)

    def tests(self):
        files=sorted(p for p in (ROOT/'src').rglob('*') if p.is_file() and p.name.startswith('test_') and p.suffix in {'.py','.cpp'})
        self.table([['Gói','Tệp test hiện có']]+[[p.relative_to(ROOT/'src').parts[0],p.name] for p in files],[.38,.62],True)

    def finish_toc(self):
        for title,anchor in self.headings:
            p=self.d.add_paragraph(style='Normal');p.paragraph_format.space_after=Pt(4)
            link=el('w:hyperlink',anchor=anchor)
            r=el('w:r');rpr=el('w:rPr');rpr.append(el('w:color',val='000000'));r.append(rpr);t=el('w:t');t.text=title;r.append(t);link.append(r);p._p.append(link)
            self.toc.addprevious(p._p)
        self.toc.getparent().remove(self.toc)

    def build(self):
        lines=(ASSET/'noi_dung.md').read_text().splitlines();i=0
        while i<len(lines):
            line=lines[i].strip();i+=1
            if not line:continue
            if line.startswith('# '):self.p(line[2:],'Title')
            elif line.startswith('## '):self.heading(line[3:],1)
            elif line.startswith('### '):self.heading(line[4:],2)
            elif line=='@TOC':
                self.p('Mục lục liên kết','Report Heading 1');self.toc=self.p()._p
            elif line=='@TRACE':self.trace()
            elif line=='@CONFIG':self.config()
            elif line=='@MANIFEST':self.manifest()
            elif line=='@TESTS':self.tests()
            elif line.startswith('$$ '):self.equation(line[3:])
            elif line.startswith('!['):
                path,caption=line[2:-1].split('|',1);self.figure(path,caption)
            elif line.startswith('|'):
                rows=[[x.strip() for x in line.strip('|').split('|')]]
                while i<len(lines) and lines[i].startswith('|'):
                    rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]);i+=1
                self.table(rows,[.24,.30,.46] if len(rows[0])==3 else None)
            else:self.p(line,'Report References' if line.startswith('[S') or line.startswith('[T') else 'Normal')
        self.finish_toc()
        settings=self.d.settings.element
        update=settings.find(qn('w:updateFields'))
        if update is None:update=el('w:updateFields',val='true');settings.append(update)
        else:update.set(qn('w:val'),'true')
        self.d.save(WORK/'working.docx')
        editable={'word/document.xml','word/_rels/document.xml.rels','word/settings.xml','docProps/core.xml','docProps/app.xml','word/styles.xml'}
        with ZipFile(REF) as src, ZipFile(WORK/'working.docx') as new, ZipFile(OUT,'w',ZIP_DEFLATED) as dst:
            for name in src.namelist():
                payload=new.read(name) if name in editable else src.read(name)
                # Add only the new Title style; preserve all other source styles exactly as elements.
                if name=='word/styles.xml':
                    oldroot=ET.fromstring(src.read(name));newroot=ET.fromstring(new.read(name))
                    replacement=newroot.find("w:style[@w:styleId='Title']",NS)
                    old=oldroot.find("w:style[@w:styleId='Title']",NS);old.getparent().replace(old,replacement)
                    payload=ET.tostring(oldroot,xml_declaration=True,encoding='UTF-8',standalone=True)
                dst.writestr(name,payload)
            for name in new.namelist():
                if name not in src.namelist():dst.writestr(name,new.read(name))
        with ZipFile(OUT) as dst,ZipFile(REF) as src:
            changed=[n for n in src.namelist() if n not in editable and src.read(n)!=dst.read(n)]
            assert not changed,changed
            assert sha(REF.read_bytes())==self.baseline['sha256']
            document=ET.fromstring(dst.read('word/document.xml'))
            # Relationship-backed images and internal links must all resolve.
            rels=ET.fromstring(dst.read('word/_rels/document.xml.rels'))
            ids={r.get('Id') for r in rels}
            for img in document.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                assert img.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') in ids
            bookmarks={e.get(qn('w:name')) for e in document.findall('.//w:bookmarkStart',NS)}
            assert all(e.get(qn('w:anchor')) in bookmarks for e in document.findall('.//w:hyperlink',NS))
        (ASSET/'figure_manifest.json').write_text(json.dumps(self.figs,ensure_ascii=False,indent=2))
        summary={'output':str(OUT),'sha256':sha(OUT.read_bytes()),'figures':len(self.figs),'equations':self.equations,'tables':self.tables,'chapters_and_appendices':len(self.headings),'preserve_only_parts_changed':changed,'source_unchanged':True,'word_field_refresh':'updateFields enabled; no Word desktop refresh performed'}
        (ASSET/'build_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
        print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':
    WORK.mkdir(exist_ok=True,parents=True);ASSET.mkdir(exist_ok=True,parents=True)
    if len(sys.argv)>1 and sys.argv[1]=='sheets':
        make_review_sheets(Path(sys.argv[2]),sys.argv[3])
    else:Builder().build()
