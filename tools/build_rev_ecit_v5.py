#!/usr/bin/env python3
"""Rebuild REV-ECIT v5 from immutable execution records and the official template.

No simulator is launched. All scientific plots use original recorded coordinates.
Dependencies: python-docx, matplotlib, numpy, scipy, Pillow, PyYAML, lxml,
latex2mathml and mathml2omml. Run with --audit-only for data checks.
"""
from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import math
import re
from collections import defaultdict
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle, FancyArrowPatch
from matplotlib.colors import ListedColormap
import numpy as np
from PIL import Image
import yaml
from scipy.stats import spearmanr
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml
import mathml2omml

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Viet_bao/bao_Rev_ecit_2026/ver5'
FIG = OUT / 'figures'
TEMPLATE = ROOT / 'Viet_bao/bao_Rev_ecit_2026/ecit_word_template.docx'
ORIGINAL = ROOT / 'Viet_bao/bao_Rev_ecit_2026/ver3/SuaREV_ECIT2026_PSTMO.docx'
CSV = ROOT / 'results/pstmo_execution_full_20260803/execution_175_cases.csv'
METHODS = ['raw', 'simple', 'savitzky_golay', 'constrained', 'pstmo']
LABEL = dict(zip(METHODS, ['Đường gốc', 'Simple', 'Savitzky–Golay', 'Constrained', 'PSTMO']))
COLORS = dict(zip(METHODS, ['#555555', '#0072B2', '#D55E00', '#9A559F', '#00866C']))
LINES = dict(zip(METHODS, [':', '--', '-.', (0, (4, 1, 1, 1)), '-']))
ENV = ['open_arena', 'research_warehouse', 'narrow_aisles', 'office_maze',
       'warehouse_cross_aisles', 'warehouse_dispatch', 'warehouse_long_aisles']
ENAMES = [f'M{i}' for i in range(1, 8)]
PLANNERS = ['NavFnAStar', 'NavFnDijkstra', 'ThetaStar', 'Smac2D', 'SmacHybrid']
CASES = [('research_warehouse', 'NavFnAStar'), ('office_maze', 'Smac2D'),
         ('warehouse_dispatch', 'Smac2D')]
SOURCE_HASHES = {}
COL_WIDTH = (11909 - 2*1361 - 360)/1440/2
FULL_WIDTH = (11909 - 2*1361)/1440


def fingerprint(path):
    path = Path(path)
    SOURCE_HASHES[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()


def load_metric_functions():
    """Extract pure numeric functions without importing ROS or changing source."""
    p = ROOT / 'src/adaptive_pivot_g2_benchmark/adaptive_pivot_g2_benchmark/compare_paths.py'
    fingerprint(p)
    module = ast.parse(p.read_text())
    names = {'calculate_path_metrics', 'resample_polyline',
             'condition_trajectory_for_metrics'}
    selected = [n for n in module.body if isinstance(n, ast.FunctionDef) and n.name in names]
    code = ast.Module(body=[ast.ImportFrom(module='__future__', names=[ast.alias(name='annotations')], level=0), *selected], type_ignores=[])
    ns = {'math': math}
    exec(compile(ast.fix_missing_locations(code), str(p), 'exec'), ns)
    return ns


def tracking_rmse(executed, reference):
    p = np.asarray(executed, dtype=float)
    r = np.asarray(reference, dtype=float)
    a, v = r[:-1], np.diff(r, axis=0)
    denom = np.sum(v*v, axis=1)
    errors = []
    for batch in np.array_split(p, max(1, len(p)//100)):
        delta = batch[:, None, :] - a[None, :, :]
        t = np.clip(np.sum(delta*v[None, :, :], axis=2)/np.maximum(denom, 1e-18), 0, 1)
        d = np.linalg.norm(delta-t[:, :, None]*v[None, :, :], axis=2).min(axis=1)
        errors.extend(d)
    return float(np.sqrt(np.mean(np.square(errors))))


def load_data():
    metric = load_metric_functions()
    for p in [CSV, TEMPLATE, ORIGINAL,
              ROOT/'results/pstmo_execution_full_20260803/execution_aggregate_5planners_7env.json']:
        fingerprint(p)
    with CSV.open(encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 175
    groups = defaultdict(dict)
    failures = []
    discrepancies = []
    for row in rows:
        p = ROOT / row['trial_json']
        fingerprint(p)
        d = json.loads(p.read_text())
        key = (row['benchmark_environment'], row['planner'])
        m = row['method']
        assert m not in groups[key]
        if row['success'] != 'True':
            failures.append({'environment':key[0], 'planner':key[1], 'method':m,
                             'source':row['trial_json'], 'error':d.get('error', d.get('controller_error_msg', ''))})
        record = {'source':row['trial_json'], 'success':row['success']=='True',
                  'raw_hash':row['raw_path_sha256']}
        for name in ['execution_time_s','tracking_rmse_m','traveled_distance_m','smoothing_time_s',
                     'planned_path_length_m','planned_max_abs_curvature_1pm','planned_curvature_energy_1pm',
                     'planned_footprint_clearance_min_m','final_position_error_m','final_yaw_error_rad',
                     'selected_path_xy','selected_path_poses','executed_path_xy','ground_truth_state_trace',
                     'configuration_sha256','ground_truth_goal_reached','physically_settled',
                     'controller_succeeded','collision_monitor_interventions',
                     'planned_footprint_collision_sample_count','start','goal']:
            record[name] = d.get(name)
        if record['success']:
            assert record['ground_truth_goal_reached'] and record['physically_settled']
            assert record['collision_monitor_interventions'] == 0
            assert record['planned_footprint_collision_sample_count'] == 0
            raw_geo = metric['calculate_path_metrics'](record['selected_path_xy'])
            for k in ['path_length_m','max_abs_curvature_1pm','curvature_energy_1pm']:
                assert math.isclose(raw_geo[k], d['planned_'+k], rel_tol=1e-9, abs_tol=1e-9)
            assert math.isclose(tracking_rmse(record['executed_path_xy'],record['selected_path_xy']),
                                d['tracking_rmse_m'],rel_tol=1e-8,abs_tol=1e-10)
            # Fixed spatial scale for comparisons; no changes to executed reference paths.
            sample = metric['resample_polyline'](record['selected_path_xy'], .05)
            geo = metric['calculate_path_metrics'](sample)
            record['geometry_uniform_005'] = geo
            record['geometry_original'] = raw_geo
            for k in ['execution_time_s','tracking_rmse_m','traveled_distance_m']:
                assert math.isclose(float(row[k]), record[k], rel_tol=1e-10)
        groups[key][m] = record
    assert len(groups)==35 and len(failures)==5
    paired = {k:v for k,v in groups.items() if set(v)==set(METHODS) and all(d['success'] for d in v.values())}
    assert len(paired)==34
    assert all(len({d['raw_hash'] for d in v.values()})==1 for v in paired.values())
    means = {}
    for m in METHODS:
        ds = [v[m] for v in paired.values()]
        means[m] = {k:float(np.mean([d[k] for d in ds])) for k in
                    ['execution_time_s','tracking_rmse_m','traveled_distance_m','smoothing_time_s','planned_path_length_m']}
        means[m].update({k:float(np.mean([d['geometry_uniform_005'][k] for d in ds]))
                         for k in ['max_abs_curvature_1pm','curvature_energy_1pm']})
    ordered = sorted(paired, key=lambda k:(ENV.index(k[0]), PLANNERS.index(k[1])))
    energy_gain = [paired[k]['simple']['geometry_uniform_005']['curvature_energy_1pm']-paired[k]['pstmo']['geometry_uniform_005']['curvature_energy_1pm'] for k in ordered]
    time_gain = [paired[k]['simple']['execution_time_s']-paired[k]['pstmo']['execution_time_s'] for k in ordered]
    rmse_gain = [paired[k]['simple']['tracking_rmse_m']-paired[k]['pstmo']['tracking_rmse_m'] for k in ordered]
    result = {'trial_count':175,'success_count':170,'paired_group_count':34,
              'geometry_spacing_m':.05,'means':means,'failures':failures,
              'pstmo_faster_than_simple':int(np.sum(np.array(time_gain)>0)),
              'pstmo_lower_rmse_than_simple':int(np.sum(np.array(rmse_gain)>0)),
              'spearman_energy_time':float(spearmanr(energy_gain,time_gain).statistic),
              'source_scalar_checks':'All successful trial RMSE and original geometry recomputed and matched',
              'cases':{f'{k[0]}/{k[1]}':{m:{x:d[x] for x in ['execution_time_s','tracking_rmse_m','traveled_distance_m','smoothing_time_s','planned_path_length_m','geometry_uniform_005','geometry_original']}
                       for m,d in paired[k].items()} for k in CASES}}
    result['all_baseline_comparisons']={}
    for m in METHODS[:-1]:
        tg=np.array([paired[k][m]['execution_time_s']-paired[k]['pstmo']['execution_time_s'] for k in ordered])
        rg=np.array([paired[k][m]['tracking_rmse_m']-paired[k]['pstmo']['tracking_rmse_m'] for k in ordered])
        eg=np.array([paired[k][m]['geometry_uniform_005']['curvature_energy_1pm']-paired[k]['pstmo']['geometry_uniform_005']['curvature_energy_1pm'] for k in ordered])
        result['all_baseline_comparisons'][m]={'time_mean_gain_s':float(tg.mean()),'time_reduction_percent':float(100*tg.mean()/means[m]['execution_time_s']),
          'time_wins':int(sum(tg>0)),'time_ties':int(sum(tg==0)),'rmse_wins':int(sum(rg>0)),'rmse_ties':int(sum(rg==0)),
          'rmse_mean_gain_mm':float(1000*rg.mean()),'spearman_energy_time':float(spearmanr(eg,tg).statistic)}
    result['failure_details']=[]
    for m in METHODS:
        rec=groups[('warehouse_dispatch','SmacHybrid')][m]
        log=ROOT/Path(rec['source']).with_suffix('.log');fingerprint(log)
        evidence=[line for line in log.read_text().splitlines() if any(x in line for x in ['Smoothed path leads to a collision','detected collision ahead','Controller patience exceeded'])]
        assert evidence
        result['failure_details'].append({'method':m,'execution_time_s':rec['execution_time_s'],'distance_m':rec['traveled_distance_m'],
          'final_error_m':rec['final_position_error_m'],'log':str(log.relative_to(ROOT)),'evidence':evidence})
    return groups, paired, ordered, result


def configure_plots():
    plt.rcParams.update({'font.family':'DejaVu Serif', 'font.size':8,
                        'axes.labelsize':8,'axes.titlesize':8,'legend.fontsize':7,
                        'xtick.labelsize':7,'ytick.labelsize':7,
                        'axes.linewidth':.55,'lines.linewidth':1.2,
                        'svg.fonttype':'none','pdf.fonttype':42})


def bezier_figure():
    """Exact quintic construction: illustrative d=1, alpha=0.30, 90 degree turn."""
    ctrl=np.array([[-1,0],[-.7,0],[-.4,0],[0,.4],[0,.7],[0,1]])
    t=np.linspace(0,1,501)
    curve=sum(math.comb(5,i)*(1-t[:,None])**(5-i)*t[:,None]**i*ctrl[i] for i in range(6))
    assert np.allclose(20*(ctrl[2]-2*ctrl[1]+ctrl[0]),0)
    assert np.allclose(20*(ctrl[5]-2*ctrl[4]+ctrl[3]),0)
    fig,ax=plt.subplots(figsize=(COL_WIDTH,2.10))
    ax.plot([-1.24,0,0],[0,0,1.20],color='#777777',ls=':',lw=1,label='Đường đa tuyến')
    ax.plot(*ctrl.T,color='#0072B2',ls='--',lw=.85,label='Đa giác điều khiển')
    ax.plot(*curve.T,color=COLORS['pstmo'],lw=1.8,label='Bézier bậc năm')
    ax.scatter(*ctrl.T,s=18,fc='white',ec='#0072B2',lw=.8,zorder=5)
    labels=[r'$A=P_0$',r'$P_1$',r'$P_2$',r'$P_3$',r'$P_4$',r'$B=P_5$']
    for i,(point,label) in enumerate(zip(ctrl,labels)):
        ax.annotate(label,point,xytext=(-1,-13) if i==0 else ((0,-13) if i<3 else (8,0)),textcoords='offset points',ha='right' if i==0 else ('center' if i<3 else 'left'),va='center',fontsize=8)
    ax.scatter(0,0,s=12,c='black');ax.text(.06,-.08,r'$V$',fontsize=8)
    for a,b,label,pos in [((-1,.19),(-.68,.19),r'$\mathbf{u}$',(-.86,.24)),((-.20,.72),(-.20,1.02),r'$\mathbf{v}$',(-.36,.82))]:
        ax.annotate('',b,xytext=a,arrowprops={'arrowstyle':'->','lw':.8,'color':'black'});ax.text(*pos,label,fontsize=8)
    for a,b in [((-1,-.34),(0,-.34)),((.64,0),(.64,1))]:
        ax.annotate('',b,xytext=a,arrowprops={'arrowstyle':'<->','lw':.6})
    ax.text(-.5,-.40,r'$d$',ha='center',fontsize=8);ax.text(.70,.5,r'$d$',fontsize=8)
    ax.text(-1.21,1.09,r'$q=\alpha d,\ \alpha=0{,}30$',fontsize=8)
    ax.set_aspect('equal');ax.set_xlim(-1.52,.90);ax.set_ylim(-.48,1.22);ax.axis('off')
    fig.legend(loc='lower center',ncol=2,frameon=False,bbox_to_anchor=(.5,0),fontsize=6.5,handlelength=1.6,columnspacing=.7)
    fig.subplots_adjust(left=.04,right=.97,top=.99,bottom=.24)
    save(fig,'figure_01_bezier')


def save(fig,name):
    fig.savefig(FIG/f'{name}.png',dpi=600,facecolor='white')
    fig.savefig(FIG/f'{name}.svg',facecolor='white')
    fig.savefig(FIG/f'{name}.pdf',facecolor='white')
    plt.close(fig)


def map_image(ax,env):
    cfg_path=ROOT/f'src/vacuum_robot_gazebo/maps/{env}.yaml'
    fingerprint(cfg_path)
    cfg=yaml.safe_load(cfg_path.read_text())
    p=cfg_path.parent/cfg['image'];fingerprint(p)
    a=np.array(Image.open(p).convert('L'))
    occupied=a<90
    x,y,_=cfg['origin'];res=cfg['resolution']
    ax.imshow(occupied,cmap=ListedColormap(['#FFFFFF','#B7B7B7']),origin='upper',
              extent=[x,x+a.shape[1]*res,y,y+a.shape[0]*res],interpolation='nearest')
    ax.set_aspect('equal',adjustable='box')
    ax.grid(color='#dedede',linewidth=.35)
    ax.set_axisbelow(True)
    ax.set_xlabel('x (m)',labelpad=1);ax.set_ylabel('y (m)',labelpad=1)
    ax.tick_params(length=2,pad=1)


def plot_paths(ax,case,paired):
    for m in METHODS:
        xy=np.asarray(paired[case][m]['selected_path_xy'])
        ax.plot(xy[:,0],xy[:,1],color=COLORS[m],ls=LINES[m],lw=1.25 if m=='pstmo' else .95,
                zorder=5 if m=='pstmo' else 3)


def case_figure(case,index,paired):
    env,planner=case
    # One overview and two independently scaled close-ups; no coordinate offsets.
    fig=plt.figure(figsize=(COL_WIDTH,3.00))
    axs=[fig.add_axes([.14,.60,.80,.34]),fig.add_axes([.13,.25,.33,.22]),fig.add_axes([.64,.25,.33,.22])]
    bounds={
      CASES[0]:[(-5.35,-1.6,-3.35,-.2),(-4.65,-3.55,-3.15,-2.15),(-3.55,-2.45,-2.75,-1.75)],
      CASES[1]:[(-5.65,5.65,-3.55,3.6),(-3.8,-2.2,-2.0,-.4),(.25,1.85,1.05,2.65)],
      CASES[2]:[(-5.65,5.65,-3.5,3.7),(-5.5,-3.8,-1.9,-.2),(2.8,4.5,-2.9,-1.2)]}[case]
    for j,(ax,b) in enumerate(zip(axs,bounds)):
        map_image(ax,env);plot_paths(ax,case,paired)
        ax.set_xlim(b[:2]);ax.set_ylim(b[2:])
        ax.set_title(['(a) Toàn tuyến','(b) Vùng 1','(c) Vùng 2'][j],pad=3,loc='left')
        if j:
            ax.xaxis.set_major_locator(plt.MaxNLocator(3));ax.yaxis.set_major_locator(plt.MaxNLocator(3))
    xy=np.asarray(paired[case]['raw']['selected_path_xy'])
    ax=axs[0]
    ax.scatter(*xy[0],marker='o',s=20,color='#ffffff',edgecolor='black',zorder=9,linewidth=.8)
    ax.scatter(*xy[-1],marker='*',s=39,color='black',zorder=9)
    for point,label,offset in [(xy[0],'Đầu',(4,-11) if case==CASES[2] else (3,6)),(xy[-1],'Đích',(-23,6) if case==CASES[2] else (-18,-12))]:
        ax.annotate(label,point,xytext=offset,textcoords='offset points',fontsize=6.5,zorder=10)
    i=len(xy)//3
    ax.annotate('',xy[min(i+7,len(xy)-1)],xytext=xy[i],arrowprops={'arrowstyle':'->','color':'black','lw':.6,'mutation_scale':7},zorder=11)
    for j,b in enumerate(bounds[1:],1):
        ax.add_patch(Rectangle((b[0],b[2]),b[1]-b[0],b[3]-b[2],fill=False,lw=.65,ls='--',ec='#222222',zorder=7))
        ax.text(b[0],b[3]+.10,str(j),fontsize=7,weight='bold')
    handles=[Line2D([0],[0],color=COLORS[m],ls=LINES[m],lw=1.3,label=LABEL[m]) for m in METHODS]
    fig.legend(handles=handles,loc='lower center',bbox_to_anchor=(.5,.005),ncol=3,
               frameon=False,handlelength=2.0,columnspacing=.7,handletextpad=.4,fontsize=6.6)
    save(fig,f'figure_{index:02d}_{env}')


def pipeline_figure():
    fig,ax=plt.subplots(figsize=(COL_WIDTH,1.73));ax.set_xlim(0,10);ax.set_ylim(0,6);ax.axis('off')
    boxes=[(0.1,4.3,2.6,1.2,'Một đường\nđầu vào'),(3.55,4.3,2.9,1.2,'Năm phương án\nxử lý đường'),
           (7.3,4.3,2.6,1.2,'Cùng bộ\nđiều khiển'),(7.3,1.9,2.6,1.2,'Robot trong\nmô phỏng'),
           (.1,.15,6.35,1.65,'Thời gian và quãng đường thực thi\nSai số bám đường\nHoàn thành nhiệm vụ')]
    for x,y,w,h,t in boxes:
        ax.add_patch(Rectangle((x,y),w,h,facecolor='white',edgecolor='black',lw=.7))
        ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=7.3)
    for a,b in [((2.7,4.9),(3.55,4.9)),((6.45,4.9),(7.3,4.9)),((8.6,4.3),(8.6,3.1)),((7.3,2.45),(6.45,.9))]:
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=8,color='black',lw=.7))
    ax.text(3.3,2.75,'Độ cong và chiều dài đường\nđược đánh giá trước thực thi',ha='center',va='center',fontsize=7.1)
    ax.add_patch(FancyArrowPatch((5,4.3),(3.3,3.35),arrowstyle='-|>',mutation_scale=8,color='black',lw=.7))
    fig.subplots_adjust(left=.01,right=.99,top=.99,bottom=.01)
    save(fig,'figure_02_protocol')


def paired_figure(paired,ordered):
    fig,axs=plt.subplots(1,2,figsize=(COL_WIDTH,2.65))
    for ax,key,scale,label in zip(axs,['execution_time_s','tracking_rmse_m'],[1,1000],['Thời gian tiết kiệm (s)','Mức giảm RMSE (mm)']):
        ax.axvline(0,c='#555555',ls='--',lw=.7)
        for env,color in zip(ENV,plt.get_cmap('tab10').colors):
            ks=[k for k in ordered if k[0]==env]
            for j,m in enumerate(METHODS[:-1]):
                xs=[(paired[k][m][key]-paired[k]['pstmo'][key])*scale for k in ks]
                ys=[j+(ENV.index(env)-3)*.06+(PLANNERS.index(k[1])-2)*.009 for k in ks]
                ax.scatter(xs,ys,s=11,c=[color],edgecolors='white',linewidths=.15,zorder=3)
        for j,m in enumerate(METHODS[:-1]):
            xs=[(paired[k][m][key]-paired[k]['pstmo'][key])*scale for k in ordered]
            ax.scatter(np.mean(xs),j,marker='|',s=120,color='black',linewidths=1.2,zorder=5)
        ax.set_yticks(range(4),[LABEL[m] for m in METHODS[:-1]] if ax is axs[0] else [])
        ax.set_ylim(3.55,-.55);ax.grid(axis='x',lw=.3,alpha=.5)
        ax.set_xlabel(label,fontsize=7,labelpad=3)
        ax.tick_params(labelsize=6.5,pad=2,length=2)
        ax.xaxis.set_major_locator(plt.MaxNLocator(4))
    hs=[Line2D([0],[0],ls='',marker='o',markersize=3.8,color=c,label=e) for e,c in zip(ENAMES,plt.get_cmap('tab10').colors)]
    fig.legend(handles=hs,loc='lower center',ncol=4,fontsize=6.5,frameon=False,columnspacing=.7)
    fig.subplots_adjust(left=.27,right=.98,top=.96,bottom=.29,wspace=.20)
    save(fig,'figure_06_paired')


def sensitivity_figure(paired,ordered):
    fig,axs=plt.subplots(2,2,figsize=(COL_WIDTH,2.65))
    for ax,m in zip(axs.flat,METHODS[:-1]):
        for env,color in zip(ENV,plt.get_cmap('tab10').colors):
            ks=[k for k in ordered if k[0]==env]
            x=[paired[k][m]['geometry_uniform_005']['curvature_energy_1pm']-paired[k]['pstmo']['geometry_uniform_005']['curvature_energy_1pm'] for k in ks]
            y=[paired[k][m]['execution_time_s']-paired[k]['pstmo']['execution_time_s'] for k in ks]
            ax.scatter(x,y,color=color,s=12,edgecolors='white',linewidths=.15)
        ax.axhline(0,color='#555555',ls='--',lw=.6);ax.axvline(0,color='#555555',ls='--',lw=.6)
        ax.set_title(LABEL[m],fontsize=7.5,pad=3);ax.grid(alpha=.3,lw=.4)
        ax.tick_params(labelsize=6.5,pad=1,length=2)
        ax.xaxis.set_major_locator(plt.MaxNLocator(3));ax.yaxis.set_major_locator(plt.MaxNLocator(3))
    fig.text(.53,.015,'Mức giảm tích phân bình phương độ cong (m⁻¹)',ha='center',fontsize=6.9)
    fig.text(.015,.54,'Thời gian tiết kiệm (s)',va='center',rotation=90,fontsize=7)
    fig.subplots_adjust(left=.19,right=.98,top=.92,bottom=.16,wspace=.35,hspace=.55)
    save(fig,'figure_07_association')


def f(value,digits=3):
    return f'{value:.{digits}f}'.replace('.',',')


class Paper:
    def __init__(self):
        self.doc=Document(TEMPLATE)
        self.body=self.doc._element.body
        self.sec=deepcopy(self.doc.sections[2]._sectPr)
        for el in list(self.body):self.body.remove(el)
        self.body.append(deepcopy(self.sec))
        self.log=[]
        for name in ['Normal','Body Text','Heading 1','Heading 2','Caption','Author','Affiliation','Abstract','key words']:
            s=self.doc.styles[name]
            s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        self.doc.styles['Normal'].font.size=Pt(10)
        s=self.doc.styles['Body Text'];s.font.size=Pt(10)
        s.paragraph_format.line_spacing=1.0;s.paragraph_format.space_after=Pt(0)
        # Number headings explicitly so citations and table numbering are deterministic.
        for name in ['Heading 1','Heading 2']:
            s=self.doc.styles[name]
            pr=s.element.get_or_add_pPr()
            for n in list(pr):
                if n.tag==qn('w:numPr'):pr.remove(n)
            s.font.size=Pt(10);s.font.bold=False
        self.doc.styles['Caption'].font.bold=False
        self.doc.styles['Caption'].font.size=Pt(8)
        self.doc.styles['Caption'].paragraph_format.space_after=Pt(5)
        self.doc.styles['Caption'].paragraph_format.line_spacing=1.0
        self.doc.core_properties.title='Đánh giá ảnh hưởng của làm mượt đường đi đến hiệu năng thực thi của robot di động vi sai'
        self.doc.core_properties.author='Hai Linh Pham; Tien Cuong Nguyen; Viet Bao Nguyen; Thi Phuong Thao Nguyen; Thi Ly Pham'
        self.doc.core_properties.subject='Đánh giá ghép cặp PSTMO trong ROS 2 Navigation2'
        self.doc.core_properties.keywords='PSTMO; robot vi sai; làm mượt đường đi; điều khiển bám đường'
        self.doc.core_properties.comments=''

    def runs(self,p,text,math_size=10):
        ppr=p._p.get_or_add_pPr()
        rpr=ppr.find(qn('w:rPr'))
        if rpr is None:rpr=OxmlElement('w:rPr');ppr.append(rpr)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(math_size*2));rpr.append(sz)
        for i,s in enumerate(re.split(r'\$(.*?)\$',text)):
            if not s:continue
            if i%2:
                mathxml=mathml2omml.convert(latex_to_mathml(s))
                wrapper=etree.fromstring(('<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'+mathxml+'</root>').encode())
                om=wrapper[0]
                ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
                for box in om.findall('.//m:box',ns):
                    bp=OxmlElement('m:boxPr');cp=OxmlElement('m:ctrlPr');rp=OxmlElement('w:rPr')
                    fs=OxmlElement('w:sz');fs.set(qn('w:val'),str(math_size*2));rp.append(fs);cp.append(rp);bp.append(cp);box.insert(0,bp)
                # Complete schema-required radical properties absent in the converter.
                for rad in om.findall('.//m:rad',ns):
                    pr=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');pr.append(hide)
                    rad.insert(0,pr);rad.insert(1,OxmlElement('m:deg'))
                for n in om.findall('.//m:nary',ns):
                    for part in ['sub','sup']:
                        if not ''.join(n.find('m:'+part,ns).itertext()):
                            hide=OxmlElement('m:'+part+'Hide');hide.set(qn('m:val'),'1');n.find('m:naryPr',ns).append(hide)
                if r'\binom' in s:
                    frac=om.find('.//m:f',ns)
                    pr=OxmlElement('m:fPr');typ=OxmlElement('m:type');typ.set(qn('m:val'),'noBar');pr.append(typ);frac.insert(0,pr)
                for mr in om.findall('.//m:r',ns):
                    pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(math_size*2));pr.append(sz)
                    sty=mr.find('m:rPr/m:sty',ns)
                    if sty is not None and sty.get(qn('m:val'))=='p':
                        italic=OxmlElement('w:i');italic.set(qn('w:val'),'0');pr.append(italic)
                    mr.insert(1 if mr.find('m:rPr',ns) is not None else 0,pr)
                p._p.append(om)
            else:p.add_run(s)

    def para(self,text,style='Body Text',size=None):
        p=self.doc.add_paragraph(style=style);self.runs(p,text)
        p.paragraph_format.widow_control=True
        if size:
            for r in p.runs:r.font.size=Pt(size)
        self.log.append(text)
        return p

    def heading(self,text,level=1):
        p=self.para(text,'Heading '+str(level))
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Pt(0)
        return p

    def equation(self,tex,number):
        p=self.doc.add_paragraph();p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(COL_WIDTH/2),1)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(COL_WIDTH),2)
        p.add_run('\t');self.runs(p,'$'+tex+'$');p.add_run('\t('+str(number)+')')
        p.paragraph_format.keep_together=True
        self.log.append(f'Equation ({number}): {tex}')

    def figure(self,name,caption):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.keep_with_next=True
        pic=p.add_run().add_picture(str(FIG/(name+'.png')),width=Inches(COL_WIDTH))
        pic._inline.docPr.set('descr',caption)
        cp=self.para(caption,'Caption');cp.paragraph_format.keep_together=True
        cp.paragraph_format.keep_with_next=False
        self.log.append(f'Figure: {name}')

    def table(self,title,headers,rows,widths=None):
        p=self.para(title,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        t=self.doc.add_table(rows=1,cols=len(headers));t.autofit=False
        t.alignment=1
        widths=widths or [COL_WIDTH/len(headers)]*len(headers)
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for i,h in enumerate(headers):
            self.runs(t.rows[0].cells[i].paragraphs[0],h,math_size=8)
        for row in rows:
            cells=t.add_row().cells
            for c,txt in zip(cells,row):c.text=str(txt)
        for ir,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ir==0:pr.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.width=Inches(widths[j]);c.vertical_alignment=1
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(1.5);pf.space_after=Pt(1.5);pf.line_spacing=1.0
                    pf.keep_with_next=ir<len(t.rows)-1
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=(ir==0)
                tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                for k in ['top','left','bottom','right']:
                    q=OxmlElement('w:'+k);q.set(qn('w:w'),'25');q.set(qn('w:type'),'dxa');mar.append(q)
                tcpr.append(mar)
        borders=OxmlElement('w:tblBorders')
        for k in ['top','bottom','insideH','left','right','insideV']:
            q=OxmlElement('w:'+k);q.set(qn('w:val'),'single' if k in ['top','bottom','insideH'] else 'nil');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'888888');borders.append(q)
        t._tbl.tblPr.append(borders)
        self.log.extend([' | '.join(headers)]+[' | '.join(map(str,r)) for r in rows])
        return t

    def title(self):
        s=self.doc.sections[0]
        cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1')
        p=self.para(self.doc.core_properties.title,'paper title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(24)
        p=self.para('Hai Linh Pham, Tien Cuong Nguyen, Viet Bao Nguyen,\nThi Phuong Thao Nguyen, Thi Ly Pham*','Author')
        p.paragraph_format.space_before=Pt(10)
        self.para('Khoa Điện – Điện tử, Trường Đại học Giao thông vận tải, Hà Nội, Việt Nam','Affiliation')
        self.para('*Tác giả liên hệ: ptyldk@utc.edu.vn','Affiliation')
        sec=self.doc.add_section(WD_SECTION_START.CONTINUOUS)
        sec._sectPr.find(qn('w:cols')).set(qn('w:num'),'2')
        sec._sectPr.find(qn('w:cols')).set(qn('w:space'),'360')

    def finish(self):
        # Match the template's final continuous section, balancing the last page.
        sec=self.doc.add_section(WD_SECTION_START.CONTINUOUS)
        sec._sectPr.find(qn('w:cols')).set(qn('w:num'),'1')
        path=OUT/'REV_ECIT2026_PSTMO_VER5.docx'
        temp=OUT/'_working.docx';self.doc.save(temp)
        # Preserve every template support part byte-for-byte unless edited content needs it.
        preserve=['word/styles.xml','word/numbering.xml','word/fontTable.xml','word/theme/theme1.xml']
        # Styles are intentionally adapted for explicit scientific heading/caption formatting.
        preserve.remove('word/styles.xml')
        with ZipFile(TEMPLATE) as a,ZipFile(temp) as b,ZipFile(path,'w',ZIP_DEFLATED) as out:
            for info in b.infolist():
                content=a.read(info.filename) if info.filename in preserve and info.filename in a.namelist() else b.read(info.filename)
                out.writestr(info,content)
        temp.unlink()
        (OUT/'manuscript.md').write_text('\n\n'.join(self.log)+'\n')
        return path


def build_document(paired,ordered,audit):
    paper=Paper();paper.title();P=paper.para;H=paper.heading;E=paper.equation;T=paper.table;F=paper.figure
    mean=audit['means'];p=mean['pstmo'];s=mean['simple']
    pct=lambda new,old:100*(old-new)/old
    P(r'Tóm tắt— Bài báo đề xuất một phương pháp hậu xử lý đường đa tuyến bằng chuyển tiếp Bézier bậc năm, được ký hiệu là PSTMO, và đánh giá ảnh hưởng của phương pháp đến thực thi của robot di động vi sai trong ROS 2 Navigation2. Đường gốc, Simple, Savitzky–Golay, Constrained và PSTMO được khảo sát qua 175 lượt trên bảy môi trường và năm bộ lập kế hoạch. So sánh ghép cặp sử dụng 34 nhóm có đủ năm lượt thành công, cùng đường đầu vào và cấu hình điều khiển. PSTMO đạt thời gian thực thi trung bình 53,148 s và căn trung bình bình phương sai số bám đường 25,51 mm, thấp nhất trong năm phương án. Thời gian giảm 2,39–7,23% so với bốn đối chứng, nhưng lợi ích không xuất hiện ở mọi nhóm và chi phí làm mượt cao hơn. Phân tích ba trường hợp cùng năm lượt thất bại chỉ ra rằng cải thiện hình học không bảo đảm giảm thời gian hoặc hoàn thành nhiệm vụ. Kết quả phản ánh cấu hình mô phỏng đã khảo sát.','Abstract')
    P(r'Từ khóa— Làm mượt đường đi; robot vi sai; PSTMO; điều khiển bám đường; Navigation2.','key words')
    H('I. GIỚI THIỆU')
    P(r'Robot di động tự hành trong kho vận và sản xuất cần thực hiện đường đi an toàn, khả thi, với thời gian phù hợp và sai số bám đường nhỏ. Trong Navigation2 (Nav2), bộ lập kế hoạch toàn cục tạo đường tham chiếu, còn bộ điều khiển sinh lệnh vận tốc từ đường và trạng thái robot [1], [2]. Vì vậy, hiệu quả của bộ làm mượt phải được đánh giá đồng thời ở cấp độ hình học và khi thực thi trong vòng điều khiển phản hồi.')
    P(r'Các phương pháp đã công bố gồm dải đàn hồi duy trì khoảng cách vật cản [3], đường giải tích liên tục độ cong [4], tối ưu gradient cho robot bánh xe [5] và các họ spline hoặc Bézier [7]–[12]. Những phương pháp này xử lý tính liên tục, giới hạn độ cong, động học và tránh vật cản theo các cách khác nhau. Các nghiên cứu về AGV còn đưa ràng buộc động học và vật cản động vào quỹ đạo spline [13]. Tuy nhiên, cải thiện một chỉ số hình học chưa tự xác định hiệu quả của đường trong vòng điều khiển kín.')
    P(r'Regulated Pure Pursuit điều chỉnh vận tốc theo độ cong và khoảng cách vật cản [6]. Bài báo này đề xuất PSTMO, một phương pháp hậu xử lý đường đa tuyến bằng chuyển tiếp Bézier bậc năm có kiểm tra hình bao và giới hạn chuyển động, sau đó đánh giá cùng ba bộ làm mượt Nav2 và đường gốc. Đóng góp gồm so sánh ghép cặp trên nhiều môi trường và bộ lập kế hoạch, đối chiếu hình học với thực thi, và phân tích cơ chế dừng của các ca thất bại. Phần II trình bày PSTMO; Phần III mô tả thử nghiệm và phân tích; Phần IV kết luận.')
    H('II. PHƯƠNG PHÁP PSTMO')
    H('A. Đường đầu vào và chuyển tiếp tại góc',2)
    P(r'PSTMO nhận đường đa tuyến từ bộ lập kế hoạch toàn cục, bản đồ chi phí, hình bao robot và các giới hạn chuyển động. Thuật toán thay cục bộ các vùng chuyển hướng, bảo toàn vị trí đầu, vị trí đích và hướng tại đích. Đầu ra là đường hình học làm tham chiếu cho bộ điều khiển; thuật toán không ấn định thời điểm robot phải đi qua từng điểm.')
    P(r'Tại đỉnh $V$, gọi $\mathbf{u}$ và $\mathbf{v}$ là vectơ đơn vị theo cạnh vào và cạnh ra. Hai điểm nối được đặt tại $A=V-d\mathbf{u}$ và $B=V+d\mathbf{v}$, với $d>0$ là chiều dài cắt trên mỗi cạnh. Đặt $q=\alpha d$, $0<\alpha\leq0,5$. Sáu điểm điều khiển được xác định bởi')
    E(r'P_0=A,\quad P_1=A+q\mathbf{u},\quad P_2=A+2q\mathbf{u}',1)
    E(r'P_3=B-2q\mathbf{v},\quad P_4=B-q\mathbf{v},\quad P_5=B',2)
    E(r'\mathbf{B}(t)=\sum_{i=0}^{5}\binom{5}{i}(1-t)^{5-i}t^iP_i,\quad t\in[0,1]',3)
    F('figure_01_bezier','Hình 1. Cấu trúc chuyển tiếp Bézier bậc năm tại góc vuông, với α = 0,30. Sáu điểm điều khiển xác định theo các Công thức 1–3; hai chiều dài cắt đều bằng d. Đây là sơ đồ hình học minh họa, không phải kết quả thực nghiệm.')
    P(r'Cách đặt điểm điều khiển khai thác tính chất cục bộ của Bézier bậc năm, một dạng cũng được dùng cho chuyển động liên tục độ cong [8]. Ba điểm ở mỗi đầu thẳng hàng và cách đều nên đạo hàm bậc hai tại hai đầu bằng không. Khi đạo hàm bậc nhất không suy biến, độ cong hai đầu bằng không, tạo mối nối $G^2$ với cạnh thẳng. Kết luận này chỉ áp dụng cho các góc được thay bằng chuyển tiếp; thao tác quay tại chỗ không được xem là một đường cong $G^2$ trên toàn miền.')
    H('B. Kiểm tra khả thi và lựa chọn chuyển tiếp',2)
    P(r'Mỗi phương án chuyển tiếp tại góc được kiểm tra theo giới hạn vận tốc tịnh tiến $v$, vận tốc góc $\omega=v\kappa$, vận tốc hai bánh và gia tốc ngang. Với khoảng cách hai bánh $b$ và độ cong $\kappa$, vận tốc bánh trái và bánh phải là')
    E(r'v_L=v(1-b\kappa/2),\quad v_R=v(1+b\kappa/2)',4)
    P(r'Đa giác hình bao robot được biến đổi theo các tư thế lấy mẫu trên từng phương án và kiểm tra với ô vật cản. Phương án không đạt điều kiện hình học, chuyển động hoặc va chạm bị loại trước bước tối ưu. Ước lượng thời gian chuyển tiếp được so sánh với phương án quay tại chỗ. Khi không có chuyển tiếp tịnh tiến phù hợp, thuật toán có thể giữ phương án quay tại chỗ nếu phương án đó vượt qua kiểm tra khả thi.')
    P(r'Hai chuyển tiếp kề nhau chỉ được ghép khi tổng chiều dài chúng sử dụng trên cạnh chung không vượt chiều dài khả dụng sau khi trừ biên phân cách. Quy hoạch động chọn chuỗi phương án tương thích có tổng chi phí nhỏ nhất; chi phí xét khoảng cách vật cản, vận tốc góc và tích phân bình phương độ cong. Sau khi ghép và lấy mẫu lại, toàn đường được kiểm tra độc lập; đường không vượt qua kiểm tra cuối bị từ chối.')
    H('III. THIẾT KẾ THỬ NGHIỆM VÀ PHÂN TÍCH')
    H('A. Hệ thống mô phỏng và phương án đối chứng',2)
    P(r'Thử nghiệm sử dụng Ubuntu 24.04, ROS 2 Jazzy, Nav2 và Gazebo Harmonic trên bộ xử lý Intel Core i5-12450HX. Robot vi sai có hình bao chữ nhật 0,44 × 0,34 m. Khoảng cách hai bánh trong mô hình kiểm tra chuyển động là 0,2548 m; mô phỏng truyền động sử dụng giá trị hiệu dụng 0,2834 m theo cấu hình đã hiệu chỉnh. Bảng I liệt kê các thông số chung.')
    T('BẢNG I. THÔNG SỐ THỬ NGHIỆM',['Đại lượng','Giá trị'],[
      ['Độ phân giải bản đồ','0,05 m/ô'],['Vận tốc tịnh tiến tối đa','0,30 m/s'],['Vận tốc góc tối đa','0,80 rad/s'],
      ['Vận tốc tuyến tính mỗi bánh tối đa','0,36 m/s'],['Gia tốc ngang tối đa','0,18 m/s²'],
      ['Bước kiểm tra hình bao / xuất đường','0,02 m / 0,05 m'],['Dung sai vị trí / hướng tại đích','0,10 m / 0,15 rad']],widths=[1.95,COL_WIDTH-1.95])
    P(r'Đường gốc giữ nguyên đường đầu vào. Simple làm trơn lặp cục bộ, tối đa 1000 vòng và hai lượt tinh chỉnh. Savitzky–Golay dùng cửa sổ bảy mẫu, hai lượt tinh chỉnh và dựa trên phép lọc bình phương tối thiểu cục bộ [14], [15]. Constrained tối ưu độ trơn và chi phí vật cản [16], với trọng số 200000 và 0,015; trọng số độ cong và độ lệch đường gốc bằng không. Vì vậy, kết quả chỉ đại diện cho cấu hình Constrained đã thử, không phải hiệu năng tốt nhất sau hiệu chỉnh tham số.')
    P(r'PSTMO giới hạn chiều dài cắt tối đa 0,80 m, phát hiện chuyển hướng từ 5° và chỉ xét chuyển tiếp tịnh tiến dưới 170°. Trọng số khoảng cách với vật cản, vận tốc góc và tích phân bình phương độ cong lần lượt là 0,15; 0,10 và 0,75. Các giới hạn trong Bảng I được giữ cố định khi thay đổi môi trường và bộ lập kế hoạch.')
    F('figure_02_protocol','Hình 2. Quy trình so sánh ghép cặp. Mỗi phương án được thực thi riêng với cùng đường đầu vào và cấu hình điều khiển.')
    P(r'Năm bộ lập kế hoạch gồm NavFn A*, NavFn Dijkstra, Theta*, Smac 2D và Smac Hybrid. Bảy môi trường được ký hiệu M1–M7 theo thứ tự cố định trong bộ dữ liệu. Mỗi môi trường sử dụng một cặp điểm đầu–đích, tạo 35 nhóm môi trường–bộ lập kế hoạch. Mỗi nhóm gồm năm phương án, tổng cộng 175 lượt. Mô phỏng được khởi tạo mới cho từng lượt; các phương án dùng cùng bộ điều khiển Regulated Pure Pursuit, cấu hình vận tốc và môi trường tĩnh. Hình 2 mô tả phép so sánh.')
    H('B. Các chỉ số đánh giá và tổng hợp',2)
    P(r'Chiều dài đường tham chiếu $L_{\mathrm{plan}}$ được tính bằng tổng khoảng cách Euclid giữa các điểm liên tiếp của đường xuất ra. Để so sánh độ cong với cùng thang không gian, tất cả đường được lấy mẫu lại theo chiều dài cung với bước 0,05 m; đoạn cuối có thể ngắn hơn. Với ba điểm liên tiếp, độ cong rời rạc được ước lượng bằng hai lần tích có hướng chia cho tích ba chiều dài cạnh tam giác. Các bộ ba suy biến có tích chiều dài không vượt $10^{-12}$ m³ được bỏ qua.')
    E(r'K_{\max}=\max_i|\kappa_i|,\quad E_\kappa=\sum_i\kappa_i^2\frac{a_i+b_i}{2}',5)
    P(r'Trong Công thức 5, $a_i$ và $b_i$ là khoảng cách từ điểm giữa tới hai điểm kề. $K_{\max}$ và $E_\kappa$ có đơn vị m⁻¹. Đây là chỉ số hình học ở bước lấy mẫu đã quy định; $E_\kappa$ không phải năng lượng điện và không đo đầy đủ chi phí quay tại chỗ. Các hình quỹ đạo hiển thị tọa độ gốc đã ghi, không dịch chuyển đường để làm nổi bật khác biệt.')
    P(r'Thời gian làm mượt $T_{\mathrm{smooth}}$ là thời lượng xử lý do máy chủ làm mượt Nav2 trả về, báo cáo bằng ms. Thời gian thực thi $T_{\mathrm{exec}}$ dùng đồng hồ mô phỏng, từ lúc bộ điều khiển nhận đường đến lúc robot dừng ổn định, gồm cả giai đoạn dừng sau khi bộ điều khiển báo hoàn tất. Điều kiện dừng yêu cầu vận tốc tịnh tiến không quá 0,01 m/s và vận tốc góc không quá 0,02 rad/s trong ba lần kiểm tra liên tiếp. Thời gian lập kế hoạch và làm mượt không nằm trong $T_{\mathrm{exec}}$.')
    P(r'Quãng đường thực thi $D_{\mathrm{exec}}$ là tổng khoảng cách giữa các vị trí thực của robot trong Gazebo; mẫu mới được lưu khi dịch chuyển ít nhất 0,1 mm. Sai số bám đường được tính từ chuỗi vị trí này sau khi lấy mẫu lại 0,05 m và lọc trung bình đối xứng tối đa năm mẫu, giữ nguyên hai đầu. Gọi $r_j^{\mathrm{f}}$ là vị trí sau xử lý và $P_m$ là đường tham chiếu của phương án $m$, ta có')
    E(r'\mathrm{RMSE}=\sqrt{\frac{1}{N}\sum_{j=1}^{N}d(r_j^{\mathrm{f}},P_m)^2}',6)
    P(r'Trong Công thức 6, $N$ là số mẫu vị trí sau xử lý; $d$ là khoảng cách nhỏ nhất tới các đoạn thẳng của đường tham chiếu. RMSE phản ánh độ bám đường riêng của mỗi phương án, không phải khoảng cách tới một đường chung và không phải sai số định vị. Phép lọc có thể làm giảm các đỉnh sai số ngắn; cùng quy trình được áp dụng cho mọi phương án.')
    P(r'Một lượt thành công khi bộ điều khiển hoàn tất, robot dừng ổn định và sai số vị trí, hướng thực tại đích nằm trong dung sai Bảng I. Tỷ lệ thành công dùng toàn bộ 35 lượt của mỗi phương án. Các giá trị trung bình và so sánh theo cặp chỉ dùng 34 nhóm có đủ năm lượt thành công và có đường đầu vào trùng khớp. Mỗi tổ hợp được chạy một lần; các chênh lệch được diễn giải dưới dạng thống kê mô tả, không dùng để khẳng định ý nghĩa thống kê.')
    H('C. Môi trường và quỹ đạo đặc trưng',2)
    H('1) Môi trường M2 với NavFn A*',2)
    P(r'Trường hợp thứ nhất sử dụng điểm đầu (−5,00; −3,00) m và điểm đích (−2,00; −0,50) m. Đường tương đối ngắn và chủ yếu chuyển hướng quanh một khối cản. Hình 3 phóng đại hai vùng khác biệt mà hình toàn tuyến khó thể hiện. Phương án PSTMO giảm dao động cục bộ trên đoạn đầu và tạo chuyển tiếp liên tục hơn khi đi vòng qua vật cản.')
    ds=paired[CASES[0]];ps,ss=ds['pstmo'],ds['simple'];eg=lambda d:d['geometry_uniform_005']['curvature_energy_1pm']
    F('figure_03_research_warehouse','Hình 3. Đường tham chiếu tại môi trường M2 với NavFn A*: (a) toàn tuyến; (b), (c) hai vùng phóng đại. Vùng xám biểu diễn vật cản; vòng tròn và ngôi sao đánh dấu đầu và đích.')
    case_table(paper,paired,CASES[0],'II')
    P(f'Bảng II cho thấy PSTMO có chiều dài {f(ps["planned_path_length_m"])} m và $E_\\kappa={f(eg(ps))}$ m⁻¹, thấp hơn Simple {f(pct(ps["planned_path_length_m"],ss["planned_path_length_m"]),2)}% và {f(pct(eg(ps),eg(ss)),2)}%. Tuy nhiên, thời gian thực thi tăng từ {f(ss["execution_time_s"])} lên {f(ps["execution_time_s"])} s, tức tăng {f(ps["execution_time_s"]-ss["execution_time_s"])} s. Ngược lại, RMSE giảm từ {f(ss["tracking_rmse_m"]*1000,2)} xuống {f(ps["tracking_rmse_m"]*1000,2)} mm. Đây là trường hợp cải thiện hình học và độ bám đường nhưng không cải thiện thời gian hoàn thành.')
    P(r'Trong cả năm phương án, PSTMO có RMSE nhỏ nhất và thời gian đứng thứ hai sau Simple. Mức giảm độ cong ở M2 vì vậy cần được tách khỏi thứ hạng thời gian. PSTMO giảm RMSE khoảng 38,18% so với Simple nhưng chậm hơn 0,222 s; mức chênh lệch nhỏ này chưa đủ để kết luận về độ ổn định khi chưa lặp phép đo. Trên tuyến ngắn, quá trình tiếp cận và dừng tại đích có thể chiếm tỷ trọng đáng kể trong thời gian đo. Tuy nhiên, dữ liệu tổng hợp chưa tách riêng thời lượng từng giai đoạn, nên đây là một khả năng giải thích, không phải nguyên nhân đã được xác nhận.')
    H('2) Môi trường M4 với Smac 2D',2)
    P(r'Trường hợp thứ hai nối (−5,20; −3,00) m với (5,20; 3,00) m qua các vách ngăn và nhiều vùng chuyển hướng. Hình 4 cho thấy các phương án đi qua cùng hành lang nhưng khác nhau tại các đoạn đổi hướng. Các vùng phóng đại giúp quan sát khác biệt cục bộ mà không thay đổi tỷ lệ hai trục tọa độ.')
    ds=paired[CASES[1]];ps,ss=ds['pstmo'],ds['simple']
    F('figure_04_office_maze','Hình 4. Đường tham chiếu tại môi trường M4 với Smac 2D. Hai khung nét đứt trong (a) xác định các vùng phóng đại (b), (c); ký hiệu và màu được giữ như Hình 3.')
    case_table(paper,paired,CASES[1],'III')
    P(f'Bảng III cho thấy, so với Simple, PSTMO giảm $E_\\kappa$ từ {f(eg(ss))} xuống {f(eg(ps))} m⁻¹, tương ứng {f(pct(eg(ps),eg(ss)),2)}%. Thời gian giảm {f(ss["execution_time_s"]-ps["execution_time_s"])} s ({f(pct(ps["execution_time_s"],ss["execution_time_s"]),2)}%), còn RMSE giảm {f(pct(ps["tracking_rmse_m"],ss["tracking_rmse_m"]),2)}%. Mặc dù PSTMO nhanh nhất trong trường hợp này, đường gốc vẫn có RMSE thấp hơn PSTMO: {f(ds["raw"]["tracking_rmse_m"]*1000,2)} so với {f(ps["tracking_rmse_m"]*1000,2)} mm. Vì vậy, thứ hạng theo thời gian và theo sai số không hoàn toàn trùng nhau.')
    P(r'PSTMO cũng nhanh hơn đường gốc, Savitzky–Golay và Constrained trong trường hợp này, nhưng RMSE còn cao hơn Constrained. Khác với M2, ưu thế hình học của PSTMO ở M4 đi cùng giảm thời gian thực thi. Kết quả phù hợp với giả thuyết rằng các vùng chuyển hướng phân bố dọc đường có thể ảnh hưởng tích lũy tới điều tiết vận tốc. Tuy nhiên, không thể quy toàn bộ mức giảm thời gian cho độ cong vì chiều dài và hình dạng đường cũng thay đổi. Mặt khác, RMSE được đo tới đường riêng của mỗi phương án: giá trị thấp hơn của đường gốc thể hiện khả năng bám đường đó, không đồng nghĩa với thời gian hoặc chất lượng hình học tốt hơn.')
    H('3) Môi trường M6 với Smac 2D',2)
    P(r'Trường hợp thứ ba đi từ (−5,20; 3,20) m đến (5,00; −2,80) m, qua hành lang bên trái rồi chuyển sang vùng phía dưới các khối cản. Hình 5 thể hiện hai vùng chuyển hướng chính. Đường dài hơn hai trường hợp trước, nhưng phần lớn chiều dài nằm trên các đoạn gần thẳng; không thể suy ra mức cải thiện chỉ từ tổng chiều dài.')
    ds=paired[CASES[2]];ps,ss=ds['pstmo'],ds['simple']
    F('figure_05_warehouse_dispatch','Hình 5. Đường tham chiếu trong môi trường M6 với Smac 2D. Các vùng phóng đại thể hiện góc rẽ ra khỏi hành lang và chuyển hướng gần đích.')
    case_table(paper,paired,CASES[2],'IV')
    P(f'Bảng IV cho thấy PSTMO có chiều dài {f(ps["planned_path_length_m"])} m, ngắn hơn Simple {f(pct(ps["planned_path_length_m"],ss["planned_path_length_m"]),2)}%. $E_\\kappa$ giảm {f(pct(eg(ps),eg(ss)),2)}%, nhưng $K_{{\\max}}$ tăng từ {f(ss["geometry_uniform_005"]["max_abs_curvature_1pm"])} lên {f(ps["geometry_uniform_005"]["max_abs_curvature_1pm"])} m⁻¹. Hai chỉ số phản ánh hai đặc điểm khác nhau: mức uốn tích lũy trên toàn đường và độ cong lớn nhất tại một vị trí. Thời gian thực thi vẫn giảm {f(ss["execution_time_s"]-ps["execution_time_s"])} s, còn RMSE giảm {f(pct(ps["tracking_rmse_m"],ss["tracking_rmse_m"]),2)}%.')
    P(r'Tại M6, PSTMO đạt thời gian và RMSE nhỏ nhất trong cả năm phương án, nhưng không đồng thời tối thiểu hóa mọi đặc trưng hình học. Một vùng cong hẹp có thể tạo giá trị cực đại lớn trong khi đóng góp vào tích phân toàn đường vẫn nhỏ. Do đó, chỉ dùng độ cong cực đại sẽ bỏ qua sự phân bố độ cong; chỉ dùng tích phân lại có thể che khuất một đoạn khó bám cục bộ. Việc thời gian giảm 1,12% dù độ cong cực đại tăng cho thấy cần đối chiếu cả hai chỉ số với kết quả điều khiển, thay vì xem một chỉ số hình học là đại diện đầy đủ cho hiệu năng thực thi.')
    H('D. Kết quả tổng hợp và so sánh ghép cặp',2)
    P(f'Bảng V sử dụng chỉ số độ cong tính lại ở cùng bước 0,05 m. PSTMO có $E_\\kappa$ trung bình {f(p["curvature_energy_1pm"])} m⁻¹, thấp hơn Simple {f(pct(p["curvature_energy_1pm"],s["curvature_energy_1pm"]),2)}%. Đổi lại, thời gian làm mượt trung bình là {f(p["smoothing_time_s"]*1000,2)} ms, cao hơn Simple ({f(s["smoothing_time_s"]*1000,2)} ms) và Constrained ({f(mean["constrained"]["smoothing_time_s"]*1000,2)} ms). Giá trị thời gian rất nhỏ chịu ảnh hưởng của độ phân giải đồng hồ; giá trị ghi bằng không ở một lượt không có nghĩa là không tốn thời gian tính toán. Đường gốc không thực hiện làm mượt nên cột này không áp dụng.')
    T('BẢNG V. KẾT QUẢ HÌNH HỌC TRUNG BÌNH TRÊN 34 NHÓM ĐẦY ĐỦ',
      ['Phương án','$L_{\\mathrm{plan}}$\n(m)','$K_{\\max}$\n(m⁻¹)','$E_\\kappa$\n(m⁻¹)','$T_{\\mathrm{smooth}}$\n(ms)'],
      [[LABEL[m],f(mean[m]['planned_path_length_m']),f(mean[m]['max_abs_curvature_1pm']),f(mean[m]['curvature_energy_1pm']),
        '—' if m=='raw' else f(mean[m]['smoothing_time_s']*1000,2)] for m in METHODS],widths=[.87,.54,.54,.54,COL_WIDTH-2.49])
    P(f'So với đường gốc, Savitzky–Golay và Constrained, PSTMO giảm $E_\\kappa$ trung bình lần lượt {f(pct(p["curvature_energy_1pm"],mean["raw"]["curvature_energy_1pm"]),2)}%, {f(pct(p["curvature_energy_1pm"],mean["savitzky_golay"]["curvature_energy_1pm"]),2)}% và {f(pct(p["curvature_energy_1pm"],mean["constrained"]["curvature_energy_1pm"]),2)}%. Độ cong cực đại và chiều dài tham chiếu trung bình cũng nhỏ nhất, trong khi chi phí làm mượt cao nhất. Đây là đánh đổi giữa mức cải thiện đường và thời gian xử lý, không phải ưu thế đồng thời trên mọi chỉ tiêu.')
    T('BẢNG VI. KẾT QUẢ THỰC THI TRUNG BÌNH TRÊN 34 NHÓM ĐẦY ĐỦ',
      ['Phương án','$D_{\\mathrm{exec}}$\n(m)','$T_{\\mathrm{exec}}$\n(s)','RMSE\n(mm)','Đạt'],
      [[LABEL[m],f(mean[m]['traveled_distance_m']),f(mean[m]['execution_time_s']),f(mean[m]['tracking_rmse_m']*1000,2),'34/35'] for m in METHODS],
      widths=[.87,.55,.58,.57,COL_WIDTH-2.57])
    c=audit['all_baseline_comparisons']
    P(r'Bảng VI cho thấy PSTMO có thời gian thực thi và RMSE trung bình thấp nhất trong năm phương án. So với đường gốc, Simple, Savitzky–Golay và Constrained, thời gian giảm lần lượt 3,973 s (6,96%), 1,301 s (2,39%), 3,529 s (6,23%) và 4,145 s (7,23%). Các tỷ lệ được tính từ chênh lệch giữa hai giá trị trung bình chia cho trung bình đối chứng, không phải trung bình của các phần trăm riêng lẻ.')
    P(r'Quãng đường thực thi trung bình của PSTMO là 10,072 m, thấp hơn Simple 0,039 m và thấp hơn ba đối chứng còn lại khoảng 0,080–0,098 m. Mức giảm nhỏ này đi cùng mức giảm thời gian khác nhau, cho thấy không thể giải thích lợi ích chỉ bằng rút ngắn đường. RMSE trung bình giảm khoảng 1,35–2,32 mm so với bốn đối chứng. Tuy nhiên, đây là mức cải thiện độ bám đường tham chiếu đã xử lý, không phải cải thiện trực tiếp độ chính xác định vị.')
    F('figure_06_paired','Hình 6. Chênh lệch đối chứng trừ PSTMO trên 34 nhóm đầy đủ cho từng đối chứng. Giá trị dương có lợi cho PSTMO. Mỗi chấm là một nhóm; vạch đen là trung bình. Màu chỉ M1–M7; độ lệch nhỏ theo phương đứng chỉ giúp phân biệt các điểm.')
    P(f'Hình 6 cho thấy PSTMO nhanh hơn đường gốc, Simple, Savitzky–Golay và Constrained lần lượt ở {c["raw"]["time_wins"]}, {c["simple"]["time_wins"]}, {c["savitzky_golay"]["time_wins"]} và {c["constrained"]["time_wins"]} trên 34 nhóm; không có trường hợp bằng nhau về thời gian. Số nhóm có RMSE nhỏ hơn tương ứng là {c["raw"]["rmse_wins"]}, {c["simple"]["rmse_wins"]}, {c["savitzky_golay"]["rmse_wins"]} và {c["constrained"]["rmse_wins"]}. Như vậy, ưu thế trung bình không có nghĩa PSTMO tốt hơn ở mọi đầu vào. Ba trường hợp trong Mục III-C minh họa những đánh đổi khác nhau; thống kê tổng hợp vẫn sử dụng đủ 34 nhóm.')
    P(f'Chi phí làm mượt cũng cần được xét cùng lợi ích thực thi. Khi cộng hai thời lượng đã ghi, PSTMO vẫn thấp hơn Simple trung bình {f((s["execution_time_s"]+s["smoothing_time_s"])-(p["execution_time_s"]+p["smoothing_time_s"]))} s. Đại lượng này không bao gồm lập kế hoạch và toàn bộ độ trễ giao tiếp; vì vậy chưa phải thời gian đáp ứng đầu cuối của hệ thống.')
    F('figure_07_association','Hình 7. Quan hệ giữa mức giảm tích phân bình phương độ cong và thời gian tiết kiệm của PSTMO so với từng đối chứng. Mỗi ô gồm 34 nhóm; giá trị dương có lợi cho PSTMO; màu như Hình 6. Các ô sử dụng thang trục riêng.')
    P(f'Hệ số tương quan hạng Spearman trong Hình 7 lần lượt là {f(c["raw"]["spearman_energy_time"])}, {f(c["simple"]["spearman_energy_time"])}, {f(c["savitzky_golay"]["spearman_energy_time"])} và {f(c["constrained"]["spearman_energy_time"])} theo thứ tự đường gốc, Simple, Savitzky–Golay và Constrained. Các hệ số được tính riêng cho từng đối chứng; không gộp 136 cặp thành các quan sát độc lập vì chúng dùng chung 34 kết quả PSTMO. Mức liên hệ phụ thuộc đối chứng và không xác lập quan hệ nhân quả. Các điểm có thời gian tiết kiệm âm dù độ cong tích lũy giảm tiếp tục cho thấy ưu thế hình học không bảo đảm ưu thế thực thi.')
    H('E. Phân tích các trường hợp thất bại',2)
    P(r'Cả năm thất bại đều thuộc M6 với Smac Hybrid, nhưng xảy ra ở hai giai đoạn khác nhau. Với Simple, máy chủ làm mượt từ chối đường do phát hiện tư thế trên đường ứng viên giao với vật cản, tại vị trí xấp xỉ (−4,741; 3,482) m. Robot chưa bước vào phép đo thực thi; vì vậy không có thời gian, quãng đường và sai số đích để so sánh như một lượt hoàn thành.')
    P(r'Với bốn phương án còn lại, bộ điều khiển liên tục cảnh báo nguy cơ va chạm phía trước, sau đó hủy thực thi khi vượt thời gian cho phép duy trì trạng thái không tạo được lệnh điều khiển hợp lệ. Bảng VII cho thấy đường gốc, Savitzky–Golay và PSTMO gần như chưa dịch chuyển, còn Constrained đi được khoảng 0,353 m. Cả bốn đều dừng ổn định nhưng còn cách đích hơn 11,7 m; dừng ổn định không đồng nghĩa với hoàn thành nhiệm vụ.')
    T('BẢNG VII. NĂM THẤT BẠI TẠI M6 VỚI SMAC HYBRID',
      ['Phương án','Giai đoạn dừng','Quãng đường (m)','Sai số đích (m)'],
      [['Đường gốc','Điều khiển','<0,001','11,834'],['Simple','Làm mượt','—','—'],
       ['Savitzky–Golay','Điều khiển','<0,001','11,834'],['Constrained','Điều khiển','0,353','11,712'],
       ['PSTMO','Điều khiển','<0,001','11,834']],widths=[.87,.84,.67,COL_WIDTH-2.38])
    P(r'Cảnh báo dự đoán của bộ điều khiển cần được phân biệt với va chạm thực tế. Bốn đường được đưa vào thực thi không có mẫu hình bao va chạm theo phép kiểm tra đã ghi và không kích hoạt bộ giám sát va chạm. Điều này không mâu thuẫn với việc bộ điều khiển từ chối chuyển động vì kiểm tra đường và dự đoán chuyển động diễn ra ở các giai đoạn khác nhau. Nhật ký xác định cơ chế dừng trực tiếp nhưng chưa tách được ảnh hưởng của bản đồ chi phí, tư thế và hướng chuyển động tới cảnh báo.')
    P(r'Thời gian hoặc RMSE rất nhỏ khi robot gần như đứng yên trong lượt bị hủy không biểu thị hiệu năng tốt. Nhóm thất bại bị loại đồng thời khỏi trung bình ghép cặp nhưng vẫn tính vào tỷ lệ thành công; mọi phương án đạt 34/35 lượt (97,14%). PSTMO chưa cải thiện tỷ lệ hoàn thành trong tập thử. Nghiên cứu chỉ chạy một lượt mỗi tổ hợp, dùng một cấu hình điều khiển và đối chứng trong mô phỏng tĩnh, nên chưa cho phép suy rộng sang độ ổn định khi lặp hoặc robot thực.')
    H('IV. KẾT LUẬN')
    P(r'Trên 34 nhóm thành công đầy đủ của 175 lượt thử, PSTMO có thời gian thực thi và RMSE trung bình thấp nhất trong năm phương án; thời gian giảm 2,39–7,23% so với bốn đối chứng. Lợi ích không đồng đều giữa các đường và đi kèm chi phí làm mượt cao hơn. Các ca thất bại cho thấy khả thi hình học và trạng thái dừng ổn định chưa bảo đảm hoàn thành nhiệm vụ. Vì vậy, đánh giá bộ làm mượt cần kết hợp hình học, thực thi và thất bại. Nghiên cứu tiếp theo cần lặp phép thử, khảo sát cấu hình và kiểm chứng trên robot thực.')
    H('TÀI LIỆU THAM KHẢO')
    refs=[
      'S. Macenski et al., “The Marathon 2: A Navigation System,” Proc. IEEE/RSJ IROS, 2020, tr. 2718–2725.',
      'S. Macenski et al., “From the Desks of ROS Maintainers: A Survey of Modern & Capable Mobile Robotics Algorithms in the Robot Operating System 2,” Robot. Auton. Syst., tập 168, mã bài 104493, 2023.',
      'S. Quinlan và O. Khatib, “Elastic Bands: Connecting Path Planning and Control,” Proc. IEEE ICRA, tập 2, 1993, tr. 802–807.',
      'K. Yang và S. Sukkarieh, “An Analytical Continuous-Curvature Path-Smoothing Algorithm,” IEEE Trans. Robot., tập 26, số 3, tr. 561–568, 2010.',
      'E. Heiden et al., “Gradient-Informed Path Smoothing for Wheeled Mobile Robots,” Proc. IEEE ICRA, 2018, tr. 1710–1717.',
      'S. Macenski et al., “Regulated Pure Pursuit for Robot Path Tracking,” Auton. Robot., tập 47, số 6, tr. 685–694, 2023.',
      'A. Ravankar et al., “Path Smoothing Techniques in Robot Navigation: State-of-the-Art, Current and Future Challenges,” Sensors, tập 18, số 9, mã bài 3170, 2018.',
      'B. Lau, C. Sprunk và W. Burgard, “Kinodynamic Motion Planning for Mobile Robots Using Splines,” Proc. IEEE/RSJ IROS, 2009, tr. 2427–2433.',
      'T. Berglund et al., “Planning Smooth and Obstacle-Avoiding B-Spline Paths for Autonomous Mining Vehicles,” IEEE Trans. Autom. Sci. Eng., tập 7, số 1, tr. 167–172, 2010.',
      'M. Elbanhawi, M. Simic và R. N. Jazar, “Continuous Path Smoothing for Car-Like Robots Using B-Spline Curves,” J. Intell. Robot. Syst., tập 80, số S1, tr. 23–56, 2015.',
      'X. Bu et al., “Curvature Continuous Path Smoothing Based on Cubic Bezier Curves for Car-Like Vehicles,” Proc. IEEE ROBIO, 2015, tr. 1453–1458.',
      'N. T. Nguyen et al., “B-Spline Path Planner for Safe Navigation of Mobile Robots,” Proc. IEEE/RSJ IROS, 2021, tr. 339–345.',
      'T. Mercy, R. Van Parys và G. Pipeleers, “Spline-Based Motion Planning for Autonomous Guided Vehicles in a Dynamic Environment,” IEEE Trans. Control Syst. Technol., tập 26, số 6, tr. 2182–2189, 2018.',
      'A. Savitzky và M. J. E. Golay, “Smoothing and Differentiation of Data by Simplified Least Squares Procedures,” Anal. Chem., tập 36, số 8, tr. 1627–1639, 1964.',
      'Navigation2 Project, “Nav2 Smoother,” tài liệu phiên bản Jazzy. https://api.nav2.org/nav2-jazzy/html/md_nav2_smoother_README.html (truy cập 14/10/2026).',
      'Navigation2 Project, “Constrained Smoother,” tài liệu phiên bản Jazzy. https://api.nav2.org/nav2-jazzy/html/md_nav2_constrained_smoother_README.html (truy cập 14/10/2026).']
    for i,r in enumerate(refs,1):
        p0=P(f'[{i}] '+r,size=8);p0.paragraph_format.first_line_indent=Inches(-.17);p0.paragraph_format.left_indent=Inches(.17);p0.paragraph_format.space_after=Pt(0)
    paper.finish()


def case_table(paper,paired,case,number):
    rows=[]
    for m in METHODS:
        d=paired[case][m]
        rows.append([LABEL[m],f(d['planned_path_length_m']),f(d['geometry_uniform_005']['curvature_energy_1pm']),f(d['execution_time_s']),f(d['tracking_rmse_m']*1000,2)])
    case_titles={CASES[0]:'MÔI TRƯỜNG M2 VỚI NAVFN A*',CASES[1]:'MÔI TRƯỜNG M4 VỚI SMAC 2D',CASES[2]:'MÔI TRƯỜNG M6 VỚI SMAC 2D'}
    paper.table(f'BẢNG {number}. '+case_titles[case],
                ['Phương án','$L_{\\mathrm{plan}}$\n(m)','$E_\\kappa$\n(m⁻¹)','$T_{\\mathrm{exec}}$\n(s)','RMSE\n(mm)'],rows,
                widths=[.87,.55,.54,.56,COL_WIDTH-2.52])


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--audit-only',action='store_true');args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True);FIG.mkdir(exist_ok=True)
    groups,paired,ordered,audit=load_data()
    (OUT/'data_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    with (OUT/'paired_measurements.csv').open('w',newline='',encoding='utf-8') as stream:
        writer=csv.writer(stream)
        writer.writerow(['environment','planner','method','reference_length_m','uniform_max_curvature_1pm',
                         'uniform_squared_curvature_integral_1pm','smoothing_duration_s','execution_time_s',
                         'executed_distance_m','tracking_rmse_m','source'])
        for env,planner in ordered:
            for method in METHODS:
                record=paired[(env,planner)][method]
                geo=record['geometry_uniform_005']
                writer.writerow([env,planner,method,record['planned_path_length_m'],geo['max_abs_curvature_1pm'],
                                 geo['curvature_energy_1pm'],record['smoothing_time_s'],record['execution_time_s'],
                                 record['traveled_distance_m'],record['tracking_rmse_m'],record['source']])
    print(f'Validated {audit["trial_count"]} trials; {audit["paired_group_count"]} complete paired groups.')
    if not args.audit_only:
        configure_plots();bezier_figure();pipeline_figure()
        for i,k in enumerate(CASES,3):case_figure(k,i,paired)
        paired_figure(paired,ordered);sensitivity_figure(paired,ordered)
        build_document(paired,ordered,audit)
    (OUT/'provenance.json').write_text(json.dumps({'sources':SOURCE_HASHES,'geometry_policy':'Uniform 0.05 m arc-length sampling of recorded reference paths; no coordinate changes to plots','tracking_policy':'Recorded conditioned trajectory: 0.05 m resampling and radius-2 moving average; nearest-segment RMSE; independently recomputed','original_template_sha256':SOURCE_HASHES[str(TEMPLATE.relative_to(ROOT))]},ensure_ascii=False,indent=2)+'\n')


if __name__=='__main__':
    main()

